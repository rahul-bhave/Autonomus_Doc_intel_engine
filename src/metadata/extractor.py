"""
Document Metadata Extractor
Pre-Sprint 3 — Client Requirement

Extracts two kinds of metadata from document files:

1. **Filesystem metadata** — size, extension, MIME type, timestamps
   Works from either a file path on disk or raw bytes + filename.

2. **Document-internal metadata** — page count, author, title, producer
   Supported formats:
     - PDF  → via pypdfium2 (already installed as docling dependency)
     - DOCX → via python-docx (if installed)
     - PPTX → via python-pptx (if installed)
"""

from __future__ import annotations

import logging
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import filetype

from src.models.schemas import DocumentMetadata

logger = logging.getLogger(__name__)


def extract_metadata(
    file_bytes: bytes,
    filename: str,
    source_path: Optional[str] = None,
) -> DocumentMetadata:
    """
    Extract filesystem + document-internal metadata.

    Args:
        file_bytes: Raw document bytes.
        filename: Original filename (used for extension detection).
        source_path: Optional absolute path on disk (enables filesystem timestamps).

    Returns:
        Populated DocumentMetadata instance.
    """
    meta = _extract_filesystem_metadata(file_bytes, filename, source_path)
    meta = _extract_document_metadata(file_bytes, meta)
    return meta


def _extract_filesystem_metadata(
    file_bytes: bytes,
    filename: str,
    source_path: Optional[str] = None,
) -> DocumentMetadata:
    """Extract size, extension, MIME type, and filesystem timestamps."""
    ext = Path(filename).suffix.lower()

    # MIME detection via magic bytes
    kind = filetype.guess(file_bytes)
    mime = kind.mime if kind else None

    # Filesystem timestamps (only available if source_path points to real file)
    created_at: Optional[datetime] = None
    modified_at: Optional[datetime] = None

    if source_path:
        try:
            stat = os.stat(source_path)
            created_at = datetime.fromtimestamp(stat.st_ctime, tz=timezone.utc)
            modified_at = datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc)
        except OSError as exc:
            logger.warning("Could not stat '%s': %s", source_path, exc)

    return DocumentMetadata(
        file_size_bytes=len(file_bytes),
        file_extension=ext,
        mime_type=mime,
        created_at=created_at,
        modified_at=modified_at,
    )


def _extract_document_metadata(
    file_bytes: bytes,
    meta: DocumentMetadata,
) -> DocumentMetadata:
    """Extract document-internal metadata based on detected type."""
    ext = meta.file_extension

    if ext == ".pdf" or (meta.mime_type and "pdf" in meta.mime_type):
        return _extract_pdf_metadata(file_bytes, meta)
    elif ext == ".docx":
        return _extract_docx_metadata(file_bytes, meta)
    elif ext == ".pptx":
        return _extract_pptx_metadata(file_bytes, meta)

    return meta


def _extract_pdf_metadata(file_bytes: bytes, meta: DocumentMetadata) -> DocumentMetadata:
    """Extract PDF metadata using pypdfium2 (already installed via docling)."""
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(file_bytes)
        meta.page_count = len(pdf)

        # pypdfium2 exposes metadata via get_metadata_dict
        try:
            pdf_meta = pdf.get_metadata_dict()
            meta.author = pdf_meta.get("Author") or None
            meta.title = pdf_meta.get("Title") or None
            meta.producer = pdf_meta.get("Producer") or None
            meta.creation_date = pdf_meta.get("CreationDate") or None
        except Exception:
            # Some PDFs have no metadata block — that's fine
            pass

        pdf.close()
    except Exception as exc:
        logger.warning("Failed to extract PDF metadata: %s", exc)

    return meta


def _extract_docx_metadata(file_bytes: bytes, meta: DocumentMetadata) -> DocumentMetadata:
    """Extract DOCX metadata using python-docx (optional dependency)."""
    try:
        from docx import Document
        from io import BytesIO

        doc = Document(BytesIO(file_bytes))
        props = doc.core_properties

        meta.author = props.author or None
        meta.title = props.title or None
        if props.created:
            meta.creation_date = props.created.isoformat()
        if props.last_modified_by:
            meta.producer = props.last_modified_by
    except ImportError:
        logger.debug("python-docx not installed — skipping DOCX internal metadata")
    except Exception as exc:
        logger.warning("Failed to extract DOCX metadata: %s", exc)

    return meta


def _extract_pptx_metadata(file_bytes: bytes, meta: DocumentMetadata) -> DocumentMetadata:
    """Extract PPTX metadata using python-pptx (optional dependency)."""
    try:
        from pptx import Presentation
        from io import BytesIO

        prs = Presentation(BytesIO(file_bytes))
        props = prs.core_properties

        meta.author = props.author or None
        meta.title = props.title or None
        if props.created:
            meta.creation_date = props.created.isoformat()
        if props.last_modified_by:
            meta.producer = props.last_modified_by

        meta.page_count = len(prs.slides)
    except ImportError:
        logger.debug("python-pptx not installed — skipping PPTX internal metadata")
    except Exception as exc:
        logger.warning("Failed to extract PPTX metadata: %s", exc)

    return meta
